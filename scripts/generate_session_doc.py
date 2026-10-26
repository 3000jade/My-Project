import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_report():
    doc = docx.Document()

    # 1. Page Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header / Footer setup
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("Human Shelter Realty  |  Session Management Technical Specification")
        f_run.font.name = "Segoe UI"
        f_run.font.size = Pt(8.5)
        f_run.font.color.rgb = RGBColor(100, 116, 139)

    # Color Palette Definitions
    COLOR_TEAL = RGBColor(15, 118, 110)       # #0F766E
    COLOR_SLATE_DARK = RGBColor(30, 41, 59)   # #1E293B
    COLOR_SLATE_MUTED = RGBColor(100, 116, 139) # #64748B
    COLOR_CORAL = RGBColor(255, 107, 107)     # #FF6B6B
    HEX_TEAL = "0F766E"
    HEX_LIGHT_TEAL = "F0FDFA"
    HEX_BG_ALT = "F8FAFC"
    HEX_BORDER = "CBD5E1"
    HEX_WHITE = "FFFFFF"

    # Helpers for Table Styling
    def set_cell_background(cell, hex_color):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_padding(cell, top=120, bottom=120, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:left w:w="{left}" w:type="dxa"/>'
            f'<w:right w:w="{right}" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)

    def set_table_borders(table, color="CBD5E1"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def style_header_cell(cell, text, width=None):
        set_cell_background(cell, HEX_TEAL)
        set_cell_padding(cell, top=140, bottom=140, left=160, right=160)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Segoe UI"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)
        if width:
            cell.width = width

    def style_data_cell(cell, text, is_even=False, is_bold=False, width=None, color=None):
        if is_even:
            set_cell_background(cell, HEX_BG_ALT)
        else:
            set_cell_background(cell, HEX_WHITE)
        set_cell_padding(cell, top=100, bottom=100, left=160, right=160)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.bold = is_bold
        run.font.name = "Segoe UI"
        run.font.size = Pt(9.0)
        run.font.color.rgb = color if color else COLOR_SLATE_DARK
        if width:
            cell.width = width

    # Heading Helpers
    def add_custom_title(title, subtitle):
        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(4)
        r_sub = p_sub.add_run("HUMAN SHELTER REALTY  |  PLATFORM ARCHITECTURE")
        r_sub.font.name = "Segoe UI"
        r_sub.font.size = Pt(10)
        r_sub.font.bold = True
        r_sub.font.color.rgb = COLOR_TEAL

        p_title = doc.add_paragraph()
        p_title.paragraph_format.space_before = Pt(4)
        p_title.paragraph_format.space_after = Pt(8)
        r_title = p_title.add_run(title)
        r_title.font.name = "Segoe UI Semibold"
        r_title.font.size = Pt(24)
        r_title.font.bold = True
        r_title.font.color.rgb = COLOR_SLATE_DARK

        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(0)
        p_desc.paragraph_format.space_after = Pt(18)
        r_desc = p_desc.add_run(subtitle)
        r_desc.font.name = "Segoe UI"
        r_desc.font.size = Pt(12)
        r_desc.font.color.rgb = COLOR_SLATE_MUTED

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Segoe UI Semibold"
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = COLOR_TEAL

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Segoe UI Semibold"
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_SLATE_DARK

    def add_body(text, space_after=6, is_italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.name = "Segoe UI"
        run.font.size = Pt(10)
        run.font.italic = is_italic
        run.font.color.rgb = COLOR_SLATE_DARK
        return p

    def add_callout(title, text, hex_bg="F0FDFA", hex_border="0F766E"):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, hex_bg)
        set_cell_padding(cell, top=120, bottom=120, left=180, right=180)
        
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{hex_border}"/>'
            f'<w:top w:val="none"/>'
            f'<w:bottom w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        r_t = p.add_run(title)
        r_t.font.name = "Segoe UI"
        r_t.font.size = Pt(10)
        r_t.font.bold = True
        r_t.font.color.rgb = COLOR_TEAL

        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.line_spacing = 1.15
        r_b = p2.add_run(text)
        r_b.font.name = "Segoe UI"
        r_b.font.size = Pt(9.5)
        r_b.font.color.rgb = COLOR_SLATE_DARK

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # DOCUMENT HEADER / TITLE
    # -------------------------------------------------------------
    add_custom_title(
        "Session Management Technical Specification & Flow Report",
        "Authoritative Dual-Layer Authentication, Inactivity Governance, Multi-Device Control, and Cross-Tab State Coordination"
    )

    # Metadata Block Table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(meta_table)
    
    meta_rows = [
        ("Platform", "Human Shelter Realty (CP_kerby Luxury Real Estate Platform)"),
        ("Architecture Tier", "Dual-Layer: Supabase Cloud Identity + Express Authoritative SessionService"),
        ("Technology Stack", "React 19 (Vite), Express 5.2 (Node.js v22), PostgreSQL with Row Level Security"),
        ("Security Standards", "SF-01 to SF-06, SESSION-01 to SESSION-03 (WCAG 2.2 AA / Zero Plaintext Tokens)"),
        ("Document Status", "Verified Production Baseline (Branch: property-tech-v3.0)"),
    ]
    for idx, (label, val) in enumerate(meta_rows):
        style_data_cell(meta_table.cell(idx, 0), label, is_even=(idx % 2 == 1), is_bold=True, width=Inches(1.8))
        style_data_cell(meta_table.cell(idx, 1), val, is_even=(idx % 2 == 1), is_bold=False, width=Inches(4.7))

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # SECTION 1: SYSTEM DESCRIPTION
    # -------------------------------------------------------------
    add_h1("1. System Description & Architectural Overview")
    
    add_body(
        "The Session Management Subsystem in Human Shelter Realty enforces an authoritative, stateful session lifecycle "
        "paired with stateless identity proofing. In luxury real estate enterprise platforms, tenant data isolation, regulatory compliance, "
        "and client confidentiality demand strict session governance that standard stateless JSON Web Tokens (JWTs) cannot provide alone."
    )

    add_body(
        "Standard JWT architectures suffer from an inherent security limitation: once signed, a bearer token remains cryptographically valid until "
        "its embedded expiration timestamp, regardless of whether a user logs out, is suspended by an administrator, or has their permissions revoked. "
        "To eliminate this vulnerability without introducing high-latency distributed token blacklists, CP_kerby implements a dual-layer security authority."
    )

    add_callout(
        "Core Dual-Layer Security Model",
        "1. Supabase Cloud Identity: Acts as the authentication provider, handling cryptographic password hashing (bcrypt), OAuth 2.0 / Google SSO token verification, and initial JWT issuance.\n"
        "2. Authoritative Express SessionService: Acts as the session governor. Every active login generates an application-level session record with strict idle and absolute time ceilings. All protected API requests pass through the requireAuth middleware, which strictly rejects any request whose session record has expired, been revoked, or is absent—regardless of JWT mathematical validity."
    )

    add_h2("1.1 Fundamental Security Principles")
    add_body(
        "• Zero Plaintext Credential Exposure at Rest (SF-01): Raw bearer tokens and JWTs are never written to disk or the database. Instead, session records are indexed via non-reversible SHA-256 cryptographic hashes (token_hash).\n"
        "• Opaque Session Identifiers: Users and client interfaces interact solely with random 36-character UUID identifiers (id). This prevents session token extraction or impersonation via client-side device inspectors.\n"
        "• Monotonic Revocation (SF-06): Session revocation is strictly unidirectional. Once a session has a non-null revoked_at timestamp, it can never be un-revoked, refreshed, or extended.\n"
        "• Server-Side Ownership Enforcement: Remote device revocation requests verify that the session's recorded userId exactly matches the authenticated caller's identity (requestingUserId === session.userId). Mismatches return 403 Forbidden.\n"
        "• Multi-Tier Storage Redundancy: Sessions are maintained in high-speed in-memory Map indices, flushed atomically to disk (backend/data/sessions.json) to survive process restarts, and synchronized with PostgreSQL (public.user_sessions) guarded by Row Level Security (RLS)."
    )

    # -------------------------------------------------------------
    # SECTION 2: FUNCTIONS & INVENTORY
    # -------------------------------------------------------------
    add_h1("2. Detailed Functions & Component Inventory")
    
    add_h2("2.1 Backend Core Services: SessionService.ts")
    add_body(
        "Location: backend/src/services/session.service.ts\n"
        "The SessionService class is the single source of truth for session lifecycles, timeouts, and multi-device state."
    )

    svc_table = doc.add_table(rows=1, cols=3)
    svc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(svc_table)
    
    style_header_cell(svc_table.cell(0, 0), "Method / Function", Inches(1.8))
    style_header_cell(svc_table.cell(0, 1), "Signature & Parameters", Inches(2.2))
    style_header_cell(svc_table.cell(0, 2), "Behavior & Guarantees", Inches(2.5))

    svc_methods = [
        ("hashToken", "hashToken(token: string): string", "Computes SHA-256 hexadecimal hash of Bearer token. Ensures zero plaintext exposure at rest."),
        ("createSession / createSessionAsync", "createSession({ sessionId?, token?, userId, role, userAgent?, ipAddress? }): AuthSessionRecord", "Creates session record with opaque UUID, calculates idleExpiresAt and absoluteExpiresAt by role, stores in RAM, writes atomic disk snapshot, and syncs to DB."),
        ("validateSession / validateSessionAsync", "validateSession(identifier, nowMs?, isHeartbeat?): { valid, reason?, session? }", "Validates presence, revocation, absolute ceiling, and idle expiry. Crucial invariant: if isHeartbeat is false, bumps idleExpiresAt. If true, leaves idle deadline untouched."),
        ("extendSession / extendSessionAsync", "extendSession(identifier, nowMs?): { valid, reason?, session? }", "Explicit server-confirmed session extension ('Stay signed in'). Re-verifies state and extends idleExpiresAt up to absolute ceiling."),
        ("rotateSession / rotateSessionAsync", "rotateSession(oldToken, newToken, userId, role, nowMs?, recordActivity?): { valid, session? }", "Updates tokenHash for refreshed token while strictly preserving original session UUID, createdAt, and absoluteExpiresAt (prevents refresh drift)."),
        ("getUserSessions / getUserSessionsAsync", "getUserSessions(userId: string, currentToken?): SessionListDTO[]", "Returns safe active session metadata for current user. Strips raw tokens and hashes; flags isCurrent for caller device."),
        ("recordHeartbeat / recordHeartbeatAsync", "recordHeartbeat(identifier, userId, nowMs?): boolean", "Validates session ownership and updates presenceStore and session.lastHeartbeatAt. Does NOT reset idle timeout."),
        ("isUserOnline", "isUserOnline(userId: string, nowMs?): boolean", "Checks whether user sent a valid heartbeat or request within the 60-second presence TTL window (PRESENCE_TTL_MS)."),
        ("revokeSession / revokeSessionAsync", "revokeSession(options, requestingUserId?, reason?, nowMs?): boolean", "Revokes target session monotonically. Enforces caller ownership check and validates permitted RevokeReason enum."),
        ("revokeAllUserSessions / revokeAllUserSessionsAsync", "revokeAllUserSessions(userId: string, reason?): number", "Terminates all active sessions across all devices for the user, clears presence, and creates an audit log entry."),
        ("disableUserAccount / disableUserAccountAsync", "disableUserAccount(userId: string, reason?): void", "Revokes all user sessions and blacklists userId in memory and persistent storage, blocking subsequent requests."),
        ("saveToPersistence / reloadFromPersistence", "saveToPersistence() / reloadFromPersistence(): void", "Atomic file serialization using temporary PID-tagged files and atomic renames to prevent partial disk writes during crashes."),
    ]

    for idx, (m_name, m_sig, m_desc) in enumerate(svc_methods):
        row = svc_table.add_row()
        is_even = (idx % 2 == 1)
        style_data_cell(row.cells[0], m_name, is_even=is_even, is_bold=True, width=Inches(1.8))
        style_data_cell(row.cells[1], m_sig, is_even=is_even, is_bold=False, width=Inches(2.2))
        style_data_cell(row.cells[2], m_desc, is_even=is_even, is_bold=False, width=Inches(2.5))

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    add_h2("2.2 Backend Controllers & Middleware")
    add_body(
        "• requireAuth (backend/src/middleware/auth.middleware.ts): Primary gatekeeper middleware. Extracts Bearer token; validates against SessionService; checks for revoked or expired state; verifies Supabase JWT cryptographic signature; confirms session-user ID match; and injects authoritative profile data into req.user.\n"
        "• optionalAuth (backend/src/middleware/auth.middleware.ts): Permissive authentication middleware. Attaches user profile if a valid token is provided, but does not block unauthenticated guest requests.\n"
        "• AuthController (backend/src/controllers/auth.controller.ts): Exposes RESTful endpoints including /api/auth/login, /api/auth/google, /api/auth/logout, /api/auth/heartbeat, /api/auth/extend-session, /api/auth/refresh, /api/auth/sessions, and /api/auth/sessions/:sessionId."
    )

    add_h2("2.3 Database Layer: public.user_sessions Schema & PostgreSQL RLS")
    add_body(
        "Location: backend/src/models/user_sessions_migration.sql\n"
        "The relational table public.user_sessions provides permanent, durable session tracking across container restarts and multi-instance cloud deployments."
    )
    add_body(
        "• Columns: id (UUID PK), user_id (UUID FK), role (TEXT), token_hash (TEXT UNIQUE), created_at, last_active_at, idle_expires_at, absolute_expires_at, revoked_at, revoke_reason, last_heartbeat_at, user_agent, ip_address, updated_at.\n"
        "• Row Level Security (RLS) Lockdown (SF-05): Direct client-side INSERT, UPDATE, DELETE permissions are revoked from authenticated and anon roles. Clients can only SELECT their own records (user_id = auth.uid()).\n"
        "• Stored Procedure public.revoke_user_session: A SECURITY DEFINER function enabling authenticated users to trigger monotonic revocation for their own sessions only."
    )

    add_h2("2.4 Frontend Services, Context & UI Components")
    add_body(
        "• sessionCoordinator.js (frontend/src/services/sessionCoordinator.js): Implements cross-tab synchronization via BroadcastChannel('cp_kerby_session_sync'). Serializes concurrent refresh token calls (executeSerializedRefresh), clears sensitive storage caches (clearSensitiveCaches), and sanitizes return paths against open redirects (validateReturnTo).\n"
        "• AuthContext.jsx (frontend/src/context/AuthContext.jsx): React context provider supplying user state, loginWithCredentials, googleLogin, logout, and token refresh listeners.\n"
        "• SessionWatcher.jsx (frontend/src/components/auth/SessionWatcher.jsx): Root-level inactivity watcher. Monitors user interaction events (mousedown, keydown, touchstart, scroll), prevents background timer drift via visibilitychange listeners, and mounts the timeout modal.\n"
        "• SessionTimeoutModal.jsx (frontend/src/components/auth/SessionTimeoutModal.jsx): Accessible countdown dialog allowing users to explicitly extend their session or log out immediately.\n"
        "• SessionManagementPanel.jsx (frontend/src/components/auth/SessionManagementPanel.jsx): Account security interface allowing users to inspect active devices, detect the current session, and remotely terminate individual browser logins."
    )

    # -------------------------------------------------------------
    # SECTION 3: CONDITIONS & POLICIES
    # -------------------------------------------------------------
    add_h1("3. Policy Conditions, Invariants & Security Rules")
    
    add_h2("3.1 Role-Based Session Policy Matrix")
    add_body("Timeout parameters are strictly differentiated based on user clearance to balance security with client convenience:")

    pol_table = doc.add_table(rows=1, cols=4)
    pol_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(pol_table)

    style_header_cell(pol_table.cell(0, 0), "Policy Dimension", Inches(1.8))
    style_header_cell(pol_table.cell(0, 1), "Staff Roles (Admin / Broker / Agent)", Inches(2.2))
    style_header_cell(pol_table.cell(0, 2), "Client Role", Inches(1.5))
    style_header_cell(pol_table.cell(0, 3), "Enforcement Layer", Inches(1.0))

    policies = [
        ("Inactivity / Idle Timeout", "30 minutes (1,800,000 ms)", "60 minutes (3,600,000 ms)", "Server-side SessionService"),
        ("Absolute Session Ceiling", "12 hours (43,200,000 ms)", "24 hours (86,400,000 ms)", "Immutable ceiling stamped at login"),
        ("Warning Lead-Time Window", "120 seconds (2 minutes)", "300 seconds (5 minutes)", "Frontend SessionWatcher modal"),
        ("Presence Heartbeat TTL", "60 seconds bounded window", "60 seconds bounded window", "SessionService presenceStore"),
        ("Heartbeat Idle Clock Reset", "Strictly Prohibited", "Strictly Prohibited", "SessionService.validateSession"),
        ("Explicit User Extension", "+30 min (capped at 12h ceiling)", "+60 min (capped at 24h ceiling)", "POST /api/auth/extend-session"),
    ]

    for idx, (p_dim, p_staff, p_client, p_layer) in enumerate(policies):
        row = pol_table.add_row()
        is_even = (idx % 2 == 1)
        style_data_cell(row.cells[0], p_dim, is_even=is_even, is_bold=True, width=Inches(1.8))
        style_data_cell(row.cells[1], p_staff, is_even=is_even, is_bold=False, width=Inches(2.2))
        style_data_cell(row.cells[2], p_client, is_even=is_even, is_bold=False, width=Inches(1.5))
        style_data_cell(row.cells[3], p_layer, is_even=is_even, is_bold=False, width=Inches(1.0))

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    add_h2("3.2 10-Step Invariant Evaluation Pipeline (requireAuth)")
    add_body("Every incoming request targeting a protected route is validated through 10 strict sequential gates:")

    pipe_table = doc.add_table(rows=1, cols=4)
    pipe_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(pipe_table)

    style_header_cell(pipe_table.cell(0, 0), "Step", Inches(0.6))
    style_header_cell(pipe_table.cell(0, 1), "Invariant Check", Inches(2.2))
    style_header_cell(pipe_table.cell(0, 2), "Failure Action & Error", Inches(2.2))
    style_header_cell(pipe_table.cell(0, 3), "Security Purpose", Inches(1.5))

    pipe_steps = [
        ("1", "Bearer Token Format Check", "401 Unauthorized ('Missing or malformed Bearer token')", "Rejects unauthenticated HTTP headers"),
        ("2", "Token Tracking Check", "401 Unauthorized ('Session not found or expired')", "Rejects un-tracked or rogue tokens"),
        ("3", "Account Disablement Check", "401 Unauthorized ('ACCOUNT_DISABLED')", "Blocks banned or suspended accounts"),
        ("4", "Revocation State Check", "401 Unauthorized ('SESSION_REVOKED')", "Enforces monotonic remote/logout terminations"),
        ("5", "Absolute Ceiling Check", "401 Unauthorized ('SESSION_ABSOLUTE_EXPIRED')", "Enforces 12h/24h non-extendable hard cap"),
        ("6", "Idle Inactivity Check", "401 Unauthorized ('SESSION_IDLE_EXPIRED')", "Terminates abandoned client sessions"),
        ("7", "Supabase Cryptographic Verification", "401 Unauthorized ('Invalid or expired session')", "Verifies digital RS256 signature"),
        ("8", "Session Ownership Invariant", "401 Unauthorized ('Session identity mismatch')", "Ensures session.userId === supabaseUser.id"),
        ("9", "Authoritative Profile Hydration", "503 Service Unavailable ('Unable to load profile')", "Loads verified role & credentials from DB"),
        ("10", "Idle Deadline Update Gate", "If isHeartbeat == true: leave deadline untouched. If false: idleExpiresAt = min(absolute, now + limit)", "Isolates presence pings from idle clocks"),
    ]

    for idx, (s_num, s_inv, s_err, s_purp) in enumerate(pipe_steps):
        row = pipe_table.add_row()
        is_even = (idx % 2 == 1)
        style_data_cell(row.cells[0], s_num, is_even=is_even, is_bold=True, width=Inches(0.6))
        style_data_cell(row.cells[1], s_inv, is_even=is_even, is_bold=False, width=Inches(2.2))
        style_data_cell(row.cells[2], s_err, is_even=is_even, is_bold=False, width=Inches(2.2))
        style_data_cell(row.cells[3], s_purp, is_even=is_even, is_bold=False, width=Inches(1.5))

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    add_h2("3.3 Sensitive Storage Sanitization Rules (SESSION-02)")
    add_body(
        "Upon logout, account switch, or role revocation, sessionCoordinator.clearSensitiveCaches() purges all keys matching prefixes: "
        "inquiry_draft_*, inquiry_cache_*, ai_summary_cache_*, chat_pending_queue_*, and broker_audit_cache_ from localStorage and sessionStorage. "
        "This guarantees client inquiry drafts and confidential valuation documents are never accessible to subsequent users on shared hardware."
    )

    add_h2("3.4 Open Redirect Defense Rules (SESSION-03)")
    add_body(
        "Return-to parameters are sanitized via validateReturnTo(rawUrl). Any path that does not begin with a single slash (/), "
        "or contains scheme delimiters (://), double slashes (//), or backslashes (/\\) is stripped and reset to root (/), "
        "preventing phishing and credential theft via malicious login redirect loops."
    )

    # -------------------------------------------------------------
    # SECTION 4: FLOWS
    # -------------------------------------------------------------
    add_h1("4. End-to-End System Flows & Lifecycles")

    flows = [
        (
            "Flow 1: Authentication & Authoritative Session Creation",
            "Covers initial user sign-in via credentials or Google OAuth SSO, issuing dual tokens and establishing persistent session tracking.",
            [
                "Step 1: User enters email/password or authenticates with Google SSO in React UI.",
                "Step 2: Frontend apiClient dispatches POST /api/auth/login or /api/auth/google.",
                "Step 3: AuthController calls Supabase Auth to authenticate credentials and issue access/refresh JWTs.",
                "Step 4: AuthController calls SessionService.createSession({ token, userId, role, userAgent, ipAddress }).",
                "Step 5: SessionService generates an opaque UUID (id), hashes the token with SHA-256 (tokenHash), and sets idleExpiresAt and absoluteExpiresAt.",
                "Step 6: SessionService writes sanitized snapshot to disk and asynchronously upserts public.user_sessions.",
                "Step 7: Server responds with HTTP 200 payload containing tokens, session ID, and user profile.",
                "Step 8: Frontend stores tokens in localStorage and redirects user to target dashboard."
            ]
        ),
        (
            "Flow 2: Authenticated API Request Validation Pipeline",
            "Covers the 10-step verification gate evaluated on every protected backend request.",
            [
                "Step 1: Browser client sends HTTP request with Authorization: Bearer <token>.",
                "Step 2: requireAuth middleware extracts Bearer token and checks SessionService.getSessionAsync(token).",
                "Step 3: If session is untracked, request is rejected immediately with HTTP 401 Unauthorized.",
                "Step 4: SessionService.validateSessionAsync verifies revocation status, absolute ceiling, and idle expiration.",
                "Step 5: If expired or revoked, request is rejected with 401 and specific reason code.",
                "Step 6: Middleware verifies cryptographic signature via supabase.auth.getUser(token).",
                "Step 7: Invariant check confirms existingSession.userId === supabaseUser.id (blocks token swapping).",
                "Step 8: Authoritative user profile loaded from PostgreSQL profiles table.",
                "Step 9: If request is not a heartbeat, idleExpiresAt is extended: min(absoluteExpiresAt, now + idleLimit).",
                "Step 10: User attached to req.user; request dispatched to target controller."
            ]
        ),
        (
            "Flow 3: Inactivity Warning & Server-Confirmed Session Extension",
            "Covers client-side idle tracking, timeout warning modal, and server-confirmed extension.",
            [
                "Step 1: SessionWatcher monitors user activity events and updates lastActivityRef.",
                "Step 2: User ceases interaction for 28 minutes (Staff role idle limit: 30 minutes).",
                "Step 3: Ticker detects remaining idle time <= 120s; mounts SessionTimeoutModal with live countdown.",
                "Step 4: User clicks 'Extend Session' button on the modal.",
                "Step 5: SessionWatcher sends POST /api/auth/extend-session with active Bearer token.",
                "Step 6: AuthController calls SessionService.extendSessionAsync(token).",
                "Step 7: Server re-verifies validity and recalculates idleExpiresAt up to absolute ceiling.",
                "Step 8: Server returns HTTP 200 with new expiration timestamp.",
                "Step 9: SessionWatcher resets local timer and broadcasts SESSION_EXTENDED via sessionCoordinator.",
                "Step 10: All other open browser tabs receive broadcast event and synchronize their local countdowns."
            ]
        ),
        (
            "Flow 4: Informational Presence Heartbeat Flow",
            "Demonstrates background presence pings updating online status without extending idle timers.",
            [
                "Step 1: Frontend background timer dispatches POST /api/auth/heartbeat with header x-session-heartbeat: true.",
                "Step 2: requireAuth recognizes heartbeat flag and passes isHeartbeat = true to SessionService.validateSession.",
                "Step 3: SessionService verifies session is valid and unrevoked, but DOES NOT modify idleExpiresAt.",
                "Step 4: SessionService updates presenceStore[userId] = nowMs and session.lastHeartbeatAt.",
                "Step 5: AuthController returns HTTP 200 { success: true, data: { online: true } }.",
                "Step 6: isUserOnline(userId) evaluates to true across the platform for next 60 seconds."
            ]
        ),
        (
            "Flow 5: Serialized Token Refresh Rotation with Ceiling Preservation",
            "Demonstrates race-free token rotation across multiple tabs while maintaining original absolute expiry.",
            [
                "Step 1: Multiple concurrent API calls or tabs encounter access token expiration.",
                "Step 2: sessionCoordinator.executeSerializedRefresh acquires activeRefreshPromise, collapsing duplicate calls.",
                "Step 3: Request sent to POST /api/auth/refresh with refreshToken and old Bearer header.",
                "Step 4: SessionService verifies old session was not revoked or disabled.",
                "Step 5: Supabase Auth exchanges refresh token and returns new access/refresh pair.",
                "Step 6: SessionService.rotateSessionAsync replaces oldTokenHash with newTokenHash in memory and database.",
                "Step 7: Original session UUID, createdAt, and absoluteExpiresAt remain strictly unmodified (SF-03).",
                "Step 8: New access token returned to client; all queued calls resolve seamlessly."
            ]
        ),
        (
            "Flow 6: Multi-Device Inspection & Remote Device Revocation",
            "Demonstrates how users view their active logins and securely terminate remote sessions.",
            [
                "Step 1: User navigates to Security / Active Sessions in React UI.",
                "Step 2: SessionManagementPanel sends GET /api/auth/sessions.",
                "Step 3: SessionService.getUserSessions filters active sessions for userId and flags isCurrent for caller.",
                "Step 4: Panel displays device list with user agent, IP address, start time, and 'Revoke' button on remote devices.",
                "Step 5: User clicks 'Revoke' on a remote session; client sends DELETE /api/auth/sessions/:sessionId.",
                "Step 6: Server verifies sessionId exists and session.userId matches requestingUserId.",
                "Step 7: SessionService stamps revokedAt = nowMs and revokeReason = 'user_manual_revoke' monotonically.",
                "Step 8: Change synced to PostgreSQL public.user_sessions table.",
                "Step 9: Next request from revoked remote device hits requireAuth, fails step 4, and receives HTTP 401 SESSION_REVOKED.",
                "Step 10: Remote client purges local tokens and redirects to login."
            ]
        ),
        (
            "Flow 7: Cross-Tab Broadcast Synchronization & Cache Purge",
            "Covers cross-tab coordination and draft data protection upon logout or account switch.",
            [
                "Step 1: User clicks 'Log Out' in Browser Tab A.",
                "Step 2: Tab A calls sessionCoordinator.broadcastSessionEvent('LOGOUT').",
                "Step 3: sessionCoordinator immediately invokes clearSensitiveCaches() in Tab A.",
                "Step 4: All inquiry draft buffers, private chat queues, and AI property summaries are purged.",
                "Step 5: BroadcastChannel('cp_kerby_session_sync') transmits event payload to all open tabs.",
                "Step 6: Tab B onmessage listener receives LOGOUT event and invokes clearSensitiveCaches().",
                "Step 7: Tab B AuthContext sets user = null; both tabs simultaneously transition to signed-out state."
            ]
        ),
        (
            "Flow 8: All-Device Revocation & Account Disablement",
            "Covers administrative account suspension and emergency all-device revocation.",
            [
                "Step 1: Administrator disables user account via Admin Console or user clicks 'Revoke All Devices'.",
                "Step 2: API dispatches call to SessionService.disableUserAccount(userId, 'account_disabled').",
                "Step 3: SessionService iterates all active sessions for userId, setting revokedAt and reason.",
                "Step 4: userId is added to disabledAccounts set and removed from presenceStore.",
                "Step 5: AuditService records auth.all_sessions_revoked event with timestamp, actor, and count.",
                "Step 6: Database public.user_sessions updated via batch update.",
                "Step 7: Any subsequent request from any device belonging to that user fails step 3 with HTTP 401 ACCOUNT_DISABLED."
            ]
        ),
    ]

    for title, desc, steps in flows:
        add_h2(title)
        add_body(desc, space_after=4, is_italic=True)
        for s in steps:
            add_body(f"• {s}", space_after=2)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # SECTION 5: QUALITY ASSURANCE & CONTROLS MATRIX
    # -------------------------------------------------------------
    add_h1("5. Quality Assurance & Security Controls Verification")
    
    add_body(
        "The session management implementation has undergone extensive automated and manual verification, "
        "covering unit test suites (session.service.test.ts), integration tests (SessionWatcher.test.jsx), "
        "and live HTTP verification against Supabase Cloud and local Express runtime."
    )

    ctrl_table = doc.add_table(rows=1, cols=3)
    ctrl_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(ctrl_table)

    style_header_cell(ctrl_table.cell(0, 0), "Control Identifier", Inches(1.8))
    style_header_cell(ctrl_table.cell(0, 1), "Security Requirement & Specification", Inches(3.2))
    style_header_cell(ctrl_table.cell(0, 2), "Verification Status", Inches(1.5))

    controls = [
        ("SF-01: Token Confidentiality", "Raw bearer tokens are never written to disk or DB; SHA-256 token_hash indexing only.", "VERIFIED & PASSING"),
        ("SF-02: Shared Authority", "PostgreSQL public.user_sessions acts as permanent durable authority with atomic disk fallback.", "VERIFIED & PASSING"),
        ("SF-03: Inactivity Isolation", "Token refresh rotations preserve original session UUID and do not extend idle expiration.", "VERIFIED & PASSING"),
        ("SF-04: Heartbeat Decoupling", "Informational presence pings update lastHeartbeatAt without postponing idle deadline.", "VERIFIED & PASSING"),
        ("SF-05: Hardened RLS", "Direct client INSERT/UPDATE/DELETE revoked; SELECT constrained to user_id = auth.uid().", "VERIFIED & PASSING"),
        ("SF-06: Strict Revocation", "Monotonic revocation with server-enforced caller ownership and strict reason enums.", "VERIFIED & PASSING"),
        ("SESSION-01: Inactivity Timers", "Staff enforced at 30 min / 12h absolute ceiling; Client enforced at 60 min / 24h ceiling.", "VERIFIED & PASSING"),
        ("SESSION-02: Cache Sanitization", "Inquiry drafts, chat queues, and AI summary caches purged on logout/account switch.", "VERIFIED & PASSING"),
        ("SESSION-03: Open Redirect Defense", "validateReturnTo rejects protocol schemes, double slashes, and backslashes.", "VERIFIED & PASSING"),
    ]

    for idx, (c_id, c_req, c_stat) in enumerate(controls):
        row = ctrl_table.add_row()
        is_even = (idx % 2 == 1)
        style_data_cell(row.cells[0], c_id, is_even=is_even, is_bold=True, width=Inches(1.8))
        style_data_cell(row.cells[1], c_req, is_even=is_even, is_bold=False, width=Inches(3.2))
        style_data_cell(row.cells[2], c_stat, is_even=is_even, is_bold=True, width=Inches(1.5), color=COLOR_TEAL)

    # Save to disk
    output_path = os.path.abspath("docs/reports/Session_Management_Technical_Specification_Report.docx")
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Report generated successfully at: {output_path}")

if __name__ == "__main__":
    create_report()
